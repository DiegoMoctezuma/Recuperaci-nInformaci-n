"""Biblioteca de layouts (todos sobre el layout en blanco, con formas y cajas posicionadas en la retícula).

Cada funcion recibe (prs, datos) y devuelve (slide, cajas). `cajas` es la lista de elementos
{tipo, x, y, w, h, texto, tamaño, color, ...} (pulgadas) que usa la vista previa y la revision.
"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from presentacion import diseno as D

P = D.PALETA
ALINEACION = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCLA = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def _rgb(h: str) -> RGBColor:
    return RGBColor(*D.rgb(h))


def _emu(v: float) -> Emu:
    return Emu(int(round(v * 914400)))


# ------------------------------------------------------------------ primitivas
def _nueva(prs, oscuro: bool):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(P["fondo_oscuro"] if oscuro else P["fondo"])
    cajas = [{"tipo": "fondo", "x": 0, "y": 0, "w": D.ANCHO, "h": D.ALTO,
              "relleno": P["fondo_oscuro"] if oscuro else P["fondo"]}]
    return slide, cajas


def _sin_estilo(shape):
    """Quita <p:style> (sombras/efectos del tema) para que solo cuente lo declarado."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _runs(texto, color):
    return list(texto) if isinstance(texto, list) else [(texto, color)]


def _poner_texto(tf, runs, pt, familia, negrita, align, ancla):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = _emu(D.INSET)
    tf.vertical_anchor = ANCLA[ancla]
    # los saltos de linea explicitos ("\n") crean parrafos nuevos
    parrafos: list[list[tuple[str, str]]] = [[]]
    for txt, col in runs:
        partes = txt.split("\n")
        for i, parte in enumerate(partes):
            if i > 0:
                parrafos.append([])
            if parte:
                parrafos[-1].append((parte, col))
    for i, pr in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALINEACION[align]
        for txt, col in pr:
            r = p.add_run()
            r.text = txt
            r.font.name = familia
            r.font.size = Pt(pt)
            r.font.bold = bool(negrita)
            r.font.color.rgb = _rgb(col)


def _texto(slide, cajas, x, y, w, h, texto, pt, color, familia=D.FUENTE_CUERPO, negrita=False,
           align="l", ancla="t", nombre="Texto"):
    tb = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tb.name = nombre
    runs = _runs(texto, color)
    _poner_texto(tb.text_frame, runs, pt, familia, negrita, align, ancla)
    cajas.append({"tipo": "texto", "x": x, "y": y, "w": w, "h": h, "texto": "".join(t for t, _ in runs),
                  "runs": runs, "tamaño": pt, "color": runs[0][1], "fuente": familia, "negrita": bool(negrita),
                  "align": align, "ancla": ancla, "nombre": nombre})
    return tb


def _rect(slide, cajas, x, y, w, h, relleno=None, borde=None, grosor=1.5, dash=False, radio=0.1,
          nombre="Forma", texto=None, pt=D.PT_ETIQUETA, color=None, tipo="rect"):
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if radio else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(forma, _emu(x), _emu(y), _emu(w), _emu(h))
    _sin_estilo(sh)
    sh.name = nombre
    if radio:
        sh.adjustments[0] = min(0.5, radio / min(w, h))
    if relleno:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(relleno)
    else:
        sh.fill.background()
    if borde:
        sh.line.color.rgb = _rgb(borde)
        sh.line.width = Pt(grosor)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    caja = {"tipo": tipo, "x": x, "y": y, "w": w, "h": h, "relleno": relleno, "borde": borde, "grosor": grosor,
            "dash": dash, "radio": radio, "nombre": nombre}
    if texto is not None:
        _poner_texto(sh.text_frame, [(texto, color)], pt, D.FUENTE_CUERPO, False, "c", "m")
        caja.update({"texto": texto, "runs": [(texto, color)], "tamaño": pt, "color": color,
                     "fuente": D.FUENTE_CUERPO, "negrita": False, "align": "c", "ancla": "m"})
    cajas.append(caja)
    return sh


def _marcador(slide, cajas, x, y, w, h, que, oscuro=False, clave=False):
    """Marcador de icono: rectangulo redondeado con borde punteado de acento y texto `ICONO: ...`."""
    color = P["neutro"] if oscuro else P["tinta_suave"]
    return _rect(slide, cajas, x, y, w, h, relleno=P["acento_suave"] if clave else None, borde=P["acento"],
                 grosor=1.75, dash=True, radio=0.15, nombre="Marcador icono", texto=f"ICONO: {que}",
                 pt=D.PT_ETIQUETA, color=P["tinta_suave"] if clave else color, tipo="marcador")


def _flecha(slide, cajas, x, y, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, _emu(x), _emu(y), _emu(w), _emu(h))
    _sin_estilo(sh)
    sh.name = "Flecha"
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(P["neutro"])
    sh.line.fill.background()
    cajas.append({"tipo": "flecha", "x": x, "y": y, "w": w, "h": h, "relleno": P["neutro"], "nombre": "Flecha"})


def _imagen(slide, cajas, ruta, x, y, w_max, h_max, alt="", h_align="l", v_align="t"):
    """Inserta la imagen conservando su proporcion dentro de la caja (w_max x h_max)."""
    ruta = Path(ruta)
    with Image.open(ruta) as im:
        iw, ih = im.size
    esc = min(w_max / iw, h_max / ih)
    w, h = iw * esc, ih * esc
    x0 = x + {"l": 0, "c": (w_max - w) / 2, "r": w_max - w}[h_align]
    y0 = y + {"t": 0, "m": (h_max - h) / 2, "b": h_max - h}[v_align]
    pic = slide.shapes.add_picture(str(ruta), _emu(x0), _emu(y0), _emu(w), _emu(h))
    pic.name = f"Imagen {ruta.stem}"
    pic._element.nvPicPr.cNvPr.set("descr", alt or ruta.stem)
    cajas.append({"tipo": "imagen", "x": x0, "y": y0, "w": w, "h": h, "ruta": str(ruta), "nombre": pic.name})
    return pic


def _pie_fijo(slide, cajas, datos):
    """Elementos fijos: nombre corto del proyecto (abajo izquierda) y numero (abajo derecha)."""
    _texto(slide, cajas, D.col_x(1), D.Y_PIE, D.span(1, 4), D.ALTO_PIE, datos["_proyecto"], D.PT_ETIQUETA,
           P["tinta_suave"], nombre="Pie proyecto")
    _texto(slide, cajas, D.col_x(11), D.Y_PIE, D.span(11, 12), D.ALTO_PIE, str(datos["_numero"]), D.PT_ETIQUETA,
           P["tinta_suave"], align="r", nombre="Numero de diapositiva")


def _titulo(slide, cajas, datos):
    """Titulo de diapositiva (frase-hallazgo). Devuelve el y donde empieza el contenido."""
    texto = datos["titulo"]
    w = D.span(1, 12)
    n = len(D.partir_lineas(texto, D.FUENTE_TITULO, True, D.PT_TITULO, (w - 2 * D.INSET) * 72))
    h = n * D.PT_TITULO * D.INTERLINEADO / 72 + 2 * D.INSET
    _texto(slide, cajas, D.col_x(1), D.Y_TITULO, w, h, texto, D.PT_TITULO, P["tinta"], D.FUENTE_TITULO, True,
           nombre="Titulo")
    return D.Y_TITULO + h + 0.25


def _h_texto(texto, pt, ancho, familia=D.FUENTE_CUERPO, negrita=False, minimo_lineas=1):
    n = max(minimo_lineas, len(D.partir_lineas(texto, familia, negrita, pt, (ancho - 2 * D.INSET) * 72)))
    return n * pt * D.INTERLINEADO / 72 + 2 * D.INSET


# ------------------------------------------------------------------ layouts
def portada(prs, d):
    slide, cajas = _nueva(prs, True)
    w = D.span(1, 9)
    resalte = d.get("titulo_resalte", "")
    t = d["titulo"]
    if resalte and t.endswith(resalte):
        runs = [(t[: -len(resalte)], P["fondo"]), (resalte, P["acento"])]
    else:
        runs = [(t, P["fondo"])]
    nt = len(D.partir_lineas(t, D.FUENTE_TITULO, True, D.PT_PORTADA, (w - 2 * D.INSET) * 72))
    ht = nt * D.PT_PORTADA * D.INTERLINEADO / 72 + 2 * D.INSET
    y = 1.5
    _texto(slide, cajas, D.col_x(1), y, w, ht, runs, D.PT_PORTADA, P["fondo"], D.FUENTE_TITULO, True, nombre="Titulo")
    hs = _h_texto(d["subtitulo"], D.PT_TEXTO, w)
    _texto(slide, cajas, D.col_x(1), y + ht + 0.15, w, hs, d["subtitulo"], D.PT_TEXTO, P["neutro"], nombre="Subtitulo")
    lado = D.span(10, 12)
    _marcador(slide, cajas, D.col_x(10), y + 0.1, lado, lado, d["icono"], oscuro=True)
    yb = D.ALTO - D.MARGEN_Y - 0.6
    _texto(slide, cajas, D.col_x(1), yb, D.span(1, 3), 0.6, d["equipo"], D.PT_ETIQUETA, P["neutro"],
           ancla="b", nombre="Equipo")
    _texto(slide, cajas, D.col_x(4), yb, D.span(4, 8), 0.6, d["materia"], D.PT_ETIQUETA, P["neutro"],
           ancla="b", nombre="Materia")
    _texto(slide, cajas, D.col_x(9), yb, D.span(9, 12), 0.6, d["fecha"], D.PT_ETIQUETA, P["neutro"],
           align="r", ancla="b", nombre="Fecha")
    return slide, cajas


def seccion(prs, d):
    slide, cajas = _nueva(prs, True)
    h = D.PT_CIFRA * D.INTERLINEADO / 72 + 2 * D.INSET
    y = (D.ALTO - h) / 2 - 0.4
    _texto(slide, cajas, D.col_x(1), y, D.span(1, 3), h, str(d["numero"]), D.PT_CIFRA, P["acento"],
           D.FUENTE_TITULO, True, ancla="m", nombre="Numero de seccion")
    _texto(slide, cajas, D.col_x(4), y, D.span(4, 12), h, d["nombre"], D.PT_PORTADA, P["fondo"],
           D.FUENTE_TITULO, True, ancla="m", nombre="Nombre de seccion")
    return slide, cajas


def cifra_grande(prs, d):
    slide, cajas = _nueva(prs, False)
    y0 = _titulo(slide, cajas, d)
    hc = D.PT_CIFRA * D.INTERLINEADO / 72 + 2 * D.INSET
    he = 0.4 if d.get("etiqueta") else 0.0
    zona = D.Y_MAX_CONTENIDO - y0
    yc = y0 + (zona - (hc + he)) / 2
    _texto(slide, cajas, D.col_x(1), yc, D.span(1, 6), hc, d["cifra"], D.PT_CIFRA, P["acento_texto"],
           D.FUENTE_TITULO, True, ancla="m", nombre="Cifra grande")
    if he:
        _texto(slide, cajas, D.col_x(1), yc + hc, D.span(1, 6), he, d["etiqueta"], D.PT_ETIQUETA, P["tinta_suave"],
               nombre="Etiqueta de la cifra")
    hf = _h_texto(d["frase"], D.PT_TEXTO, D.span(7, 12))
    _texto(slide, cajas, D.col_x(7), y0 + (zona - hf) / 2, D.span(7, 12), hf, d["frase"], D.PT_TEXTO, P["tinta"],
           ancla="m", nombre="Frase de apoyo")
    _pie_fijo(slide, cajas, d)
    return slide, cajas


def grafica(prs, d):
    slide, cajas = _nueva(prs, False)
    y0 = _titulo(slide, cajas, d)
    _imagen(slide, cajas, d["imagen"], D.col_x(1), y0, D.span(1, 8), D.Y_MAX_CONTENIDO - y0, d.get("alt", ""))
    hm = _h_texto(d["mensaje"], D.PT_TEXTO, D.span(9, 12))
    _texto(slide, cajas, D.col_x(9), y0, D.span(9, 12), hm, d["mensaje"], D.PT_TEXTO, P["tinta"], nombre="Mensaje clave")
    if d.get("pie"):
        hp = _h_texto(d["pie"], D.PT_ETIQUETA, D.span(9, 12))
        _texto(slide, cajas, D.col_x(9), D.Y_MAX_CONTENIDO - hp, D.span(9, 12), hp, d["pie"], D.PT_ETIQUETA,
               P["tinta_suave"], ancla="b", nombre="Pie de fuente")
    _pie_fijo(slide, cajas, d)
    return slide, cajas


def grafica_completa(prs, d):
    slide, cajas = _nueva(prs, False)
    y0 = _titulo(slide, cajas, d)
    hp = _h_texto(d["pie"], D.PT_ETIQUETA, D.span(1, 12))
    ypie = D.Y_MAX_CONTENIDO - hp
    _imagen(slide, cajas, d["imagen"], D.col_x(1), y0, D.span(1, 12), ypie - 0.05 - y0, d.get("alt", ""))
    _texto(slide, cajas, D.col_x(1), ypie, D.span(1, 12), hp, d["pie"], D.PT_ETIQUETA, P["tinta_suave"],
           nombre="Pie de fuente")
    _pie_fijo(slide, cajas, d)
    return slide, cajas


def bloques(prs, d):
    slide, cajas = _nueva(prs, False)
    y0 = _titulo(slide, cajas, d)
    bl = d["bloques"]
    n = len(bl)
    assert 2 <= n <= 4, "bloques admite de 2 a 4 bloques"
    cols = {2: 6, 3: 4, 4: 3}[n]
    w = D.span(1, cols)
    # alto del visual: imagen (si hay) o marcador de icono
    hv = 2.3 if any("imagen" in b for b in bl) else 1.7
    hn = 20 * D.INTERLINEADO / 72 + 2 * D.INSET
    hf = max(_h_texto(b["frase"], D.PT_TEXTO, w) for b in bl)
    hp = _h_texto(d["pie"], D.PT_ETIQUETA, D.span(1, 12)) if d.get("pie") else 0.0
    total = hv + 0.1 + hn + hf
    ytop = y0 + max(0.0, (D.Y_MAX_CONTENIDO - hp - y0 - total) / 2)
    for i, b in enumerate(bl):
        x = D.col_x(1 + i * cols)
        if "imagen" in b:
            _imagen(slide, cajas, b["imagen"], x + D.INSET, ytop, w - D.INSET - 0.1, hv, b.get("alt", b["nombre"]), "l", "t")
        else:
            _marcador(slide, cajas, x + D.INSET, ytop, hv, hv, b["icono"])
        _texto(slide, cajas, x, ytop + hv + 0.1, w, hn, b["nombre"], D.PT_TEXTO, P["tinta"], negrita=True,
               nombre=f"Bloque {i + 1} nombre")
        _texto(slide, cajas, x, ytop + hv + 0.1 + hn, w, hf, b["frase"], D.PT_TEXTO, P["tinta"],
               nombre=f"Bloque {i + 1} frase")
    if hp:
        _texto(slide, cajas, D.col_x(1), D.Y_MAX_CONTENIDO - hp, D.span(1, 12), hp, d["pie"], D.PT_ETIQUETA,
               P["tinta_suave"], nombre="Pie de fuente")
    _pie_fijo(slide, cajas, d)
    return slide, cajas


def flujo(prs, d):
    slide, cajas = _nueva(prs, False)
    y0 = _titulo(slide, cajas, d)
    pasos = d["pasos"]
    n = len(pasos)
    hueco = 0.5
    w = (D.span(1, 12) - (n - 1) * hueco) / n
    lado = 1.8
    hl = 20 * D.INTERLINEADO / 72 + 2 * D.INSET
    hd = max(_h_texto(p.get("detalle", ""), D.PT_ETIQUETA, w) for p in pasos)
    total = lado + 0.15 + hl + hd
    ytop = y0 + max(0.0, (D.Y_MAX_CONTENIDO - y0 - total) / 2)
    for i, p in enumerate(pasos):
        x = D.col_x(1) + i * (w + hueco)
        clave = bool(p.get("clave"))
        _marcador(slide, cajas, x + (w - lado) / 2, ytop, lado, lado, p["icono"], clave=clave)
        _texto(slide, cajas, x, ytop + lado + 0.15, w, hl, p["etiqueta"], D.PT_TEXTO,
               P["acento_texto"] if clave else P["tinta"], negrita=True, align="c", nombre=f"Paso {i + 1} etiqueta")
        _texto(slide, cajas, x, ytop + lado + 0.15 + hl, w, hd, p.get("detalle", ""), D.PT_ETIQUETA, P["tinta_suave"],
               align="c", nombre=f"Paso {i + 1} detalle")
        if i < n - 1:
            aw, ah = 0.42, 0.36
            _flecha(slide, cajas, x + w + (hueco - aw) / 2, ytop + (lado - ah) / 2, aw, ah)
    _pie_fijo(slide, cajas, d)
    return slide, cajas


def demo(prs, d):
    slide, cajas = _nueva(prs, False)
    y0 = _titulo(slide, cajas, d)
    hb = 0.7
    X0, WW = D.col_x(1) + D.INSET, D.span(1, 12) - 2 * D.INSET   # alineado con el texto del titulo
    _rect(slide, cajas, X0, y0, WW, hb, relleno=P["superficie"], borde=P["tinta_suave"],
          grosor=1.5, radio=0.35, nombre="Barra de busqueda")
    _texto(slide, cajas, X0 + 0.3, y0, WW - 0.6, hb, d["consulta"], 22, P["tinta"],
           ancla="m", nombre="Consulta")
    hc, sep = 1.05, 0.15
    y = y0 + hb + 0.25
    for i, r in enumerate(d["resultados"][:3]):
        _rect(slide, cajas, X0, y, WW, hc, relleno=P["superficie"], borde=P["neutro"],
              grosor=1.0, radio=0.12, nombre=f"Tarjeta {i + 1}")
        _texto(slide, cajas, X0, y, D.span(1, 1), hc, str(r["rank"]), D.PT_TITULO,
               P["acento_texto"] if i == 0 else P["tinta_suave"], D.FUENTE_TITULO, True, align="c", ancla="m",
               nombre=f"Rango {i + 1}")
        xb = D.col_x(2) + D.INSET
        cat = r["category"]
        _rect(slide, cajas, xb + D.INSET, y + 0.05 + 0.11, 0.18, 0.18, relleno=D.CATEGORIAS[cat], radio=0,
              nombre=f"Color categoria {i + 1}")
        _texto(slide, cajas, xb + 0.3, y + 0.05, D.span(2, 12) - 0.4, 0.4, D.NOMBRE_CATEGORIA[cat], D.PT_ETIQUETA,
               P["tinta_suave"], ancla="m", nombre=f"Categoria {i + 1}")
        _texto(slide, cajas, xb, y + 0.45, D.span(2, 12) - 0.2, 0.55, r["titulo"], D.PT_TEXTO, P["tinta"],
               nombre=f"Titulo resultado {i + 1}")
        y += hc + sep
    _pie_fijo(slide, cajas, d)
    return slide, cajas


def cierre(prs, d):
    slide, cajas = _nueva(prs, True)
    w = D.span(1, 10)
    hc = _h_texto(d["conclusion"], D.PT_TITULO, w, D.FUENTE_TITULO, True)
    hr = _h_texto(d["recomendacion"], D.PT_TEXTO, D.span(1, 9))
    total = hc + 0.6 + 0.4 + hr
    y = (D.ALTO - total) / 2
    _texto(slide, cajas, D.col_x(1), y, w, hc, d["conclusion"], D.PT_TITULO, P["fondo"], D.FUENTE_TITULO, True,
           nombre="Conclusion")
    y += hc + 0.6
    _texto(slide, cajas, D.col_x(1), y, D.span(1, 9), 0.4, d.get("etiqueta", "Recomendación"), D.PT_ETIQUETA,
           P["acento"], negrita=True, nombre="Etiqueta recomendacion")
    _texto(slide, cajas, D.col_x(1), y + 0.4, D.span(1, 9), hr, d["recomendacion"], D.PT_TEXTO, P["fondo"],
           nombre="Recomendacion")
    return slide, cajas


LAYOUT_FUNCIONES = {f.__name__: f for f in (portada, seccion, cifra_grande, grafica, grafica_completa, bloques,
                                            flujo, demo, cierre)}
