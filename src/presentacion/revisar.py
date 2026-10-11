"""Revision automatica del .pptx final (paso 6A de la skill).

Lee presentacion/salida/presentacion_final.pptx (formas, textos, colores) y escribe
presentacion/salida/revision_automatica.json con {"errores", "avisos", "detalle"}.
Uso (desde la raiz):  conda run -n bbc-ir python src/presentacion/revisar.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from PIL import Image

from presentacion import diseno as D

EMU = 914400
TOL = 0.002           # pulgadas
ACENTOS = {D.PALETA["acento"].upper(), D.PALETA["acento_texto"].upper()}
RUTA_JSON = D.DIR_SALIDA / "revision_automatica.json"


def _pulg(v) -> float:
    return v / EMU


def _color_run(run, defecto):
    try:
        return "#" + str(run.font.color.rgb)
    except Exception:  # noqa: BLE001
        return defecto


def _fondo_slide(slide) -> str:
    try:
        return "#" + str(slide.background.fill.fore_color.rgb)
    except Exception:  # noqa: BLE001
        return D.PALETA["fondo"]


def _relleno(sh):
    """Color de relleno solido de una autoforma (o None)."""
    try:
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.fill.type == 1:
            return "#" + str(sh.fill.fore_color.rgb)
    except Exception:  # noqa: BLE001
        pass
    return None


def revisar() -> dict:
    prs = Presentation(D.RUTA_PPTX)
    W, H = _pulg(prs.slide_width), _pulg(prs.slide_height)
    det: list[dict] = []
    errores = avisos = 0

    def reg(nivel, n, regla, msg):
        nonlocal errores, avisos
        if nivel == "error":
            errores += 1
        else:
            avisos += 1
        det.append({"nivel": nivel, "diapositiva": n, "regla": regla, "mensaje": msg})

    resumen_palabras = {}
    for n, slide in enumerate(prs.slides, 1):
        fondo = _fondo_slide(slide)
        formas = list(slide.shapes)
        textos = []          # (shape, x, y, w, h)
        palabras = 0
        acento_elems = 0

        # sin animaciones ni transiciones
        xml = slide._element
        if xml.find(qn("p:transition")) is not None or xml.find(qn("p:timing")) is not None:
            reg("error", n, "animaciones", "la diapositiva tiene transicion o animacion")

        for idx, sh in enumerate(formas):
            x, y, w, h = (_pulg(v) for v in (sh.left, sh.top, sh.width, sh.height))
            # 1) limites: lienzo (error) y margenes (error)
            if x < -TOL or y < -TOL or x + w > W + TOL or y + h > H + TOL:
                reg("error", n, "limites", f"'{sh.name}' sale del lienzo")
            elif (x < D.MARGEN_X - TOL or x + w > W - D.MARGEN_X + TOL or
                  y < D.MARGEN_Y - TOL or y + h > H - D.MARGEN_Y + TOL):
                # la portada/cierre usan el mismo margen; solo se tolera el lienzo completo (fondo no es forma)
                reg("error", n, "margenes", f"'{sh.name}' invade el margen del lienzo ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")

            # 2) imagenes: proporcion
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                iw, ih = Image.open(__import__("io").BytesIO(sh.image.blob)).size
                dif = abs((w / h) / (iw / ih) - 1)
                if dif >= 0.01:
                    reg("error", n, "proporcion", f"'{sh.name}' deforma la imagen ({dif * 100:.2f}%)")
                continue

            if not sh.has_text_frame:
                relleno = _relleno(sh)
                if relleno and relleno.upper() in ACENTOS and not sh.name.startswith("Marcador"):
                    acento_elems += 1
                continue
            tf = sh.text_frame
            txt = tf.text.strip()
            if not txt:
                continue
            palabras += len(txt.split())
            textos.append((sh, x, y, w, h))

            # 3) formato: viñetas, niveles, tamaños, fuentes
            ancho_pt = (w - 2 * D.INSET) * 72
            alto_pt = (h - 2 * D.INSET) * 72
            total_lineas, alto_necesario = 0, 0.0
            acento_texto = False
            bg = None
            # fondo bajo el texto: ultimo relleno solido anterior (en z-order) que contiene el centro, o el propio
            cx, cy = x + w / 2, y + h / 2
            bg = _relleno(sh)
            if bg is None:
                for otro in formas[:idx]:
                    r = _relleno(otro)
                    if r and _pulg(otro.left) <= cx <= _pulg(otro.left + otro.width) and _pulg(otro.top) <= cy <= _pulg(otro.top + otro.height):
                        bg = r
            bg = bg or fondo
            for p in tf.paragraphs:
                texto_p = "".join(r.text for r in p.runs)
                if not texto_p:
                    continue
                if p.level and p.level > 0:
                    reg("error", n, "nivel", f"'{sh.name}': parrafo con level > 0")
                if texto_p.lstrip()[:1] in ("•", "-", "*", "–"):
                    reg("error", n, "vineta", f"'{sh.name}': parrafo empieza con '{texto_p[:1]}'")
                ppr = p._p.find(qn("a:pPr"))
                if ppr is not None and any(ppr.find(qn(t)) is not None for t in ("a:buChar", "a:buAutoNum", "a:buBlip")):
                    reg("error", n, "vineta", f"'{sh.name}': parrafo con vineta")
                runs = [r for r in p.runs if r.text]
                pt = max(r.font.size.pt for r in runs)
                fam = runs[0].font.name
                neg = bool(runs[0].font.bold)
                if min(r.font.size.pt for r in runs) < 12:
                    reg("error", n, "tamano", f"'{sh.name}': texto menor a 12 pt")
                if fam not in (D.FUENTE_TITULO, D.FUENTE_CUERPO):
                    reg("error", n, "fuente", f"'{sh.name}': fuente '{fam}' fuera del sistema")
                lineas = D.partir_lineas(texto_p, fam, neg, pt, ancho_pt)
                total_lineas += len(lineas)
                alto_necesario += len(lineas) * pt * D.INTERLINEADO
                pal_max = D.max_palabra(texto_p, fam, neg, pt)
                if pal_max > ancho_pt:
                    reg("error", n, "ajuste", f"'{sh.name}': una palabra ({pal_max / 72:.2f} in) no cabe en el ancho ({ancho_pt / 72:.2f} in)")
                for r in runs:
                    c = _color_run(r, D.PALETA["tinta"])
                    cr = D.contraste(c, bg)
                    if cr < 4.5:
                        reg("error", n, "contraste", f"'{sh.name}': contraste {cr:.2f}:1 de {c} sobre {bg}")
                    if c.upper() in ACENTOS:
                        acento_texto = True
            if alto_necesario > alto_pt + 0.5:
                reg("error", n, "ajuste", f"'{sh.name}': el texto necesita {alto_necesario / 72:.2f} in y la caja ofrece {alto_pt / 72:.2f} in ({total_lineas} lineas)")
            if acento_texto and not sh.name.startswith("Marcador"):
                acento_elems += 1

        # 4) solapes entre cajas de texto (los marcadores son cajas de texto tambien)
        for i in range(len(textos)):
            for j in range(i + 1, len(textos)):
                a, b = textos[i], textos[j]
                ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
                if ox > TOL and oy > TOL:
                    reg("error", n, "solape", f"'{a[0].name}' se solapa con '{b[0].name}'")

        # 5) palabras
        resumen_palabras[n] = palabras
        if palabras > 60:
            reg("error", n, "palabras", f"{palabras} palabras (maximo 60)")
        elif palabras > 55:
            reg("aviso", n, "palabras", f"{palabras} palabras: cerca del limite de 60")

        # 6) un solo acento protagonista
        if acento_elems > 1:
            reg("aviso", n, "acento", f"{acento_elems} elementos en color acento (se recomienda uno)")

        # 7) notas del orador
        nota = slide.notes_slide.notes_text_frame.text.split() if slide.has_notes_slide else []
        if not 40 <= len(nota) <= 80:
            reg("error", n, "notas", f"notas con {len(nota)} palabras (40-80)")

    out = {"errores": errores, "avisos": avisos, "detalle": det,
           "palabras_por_diapositiva": resumen_palabras, "n_diapositivas": len(prs.slides),
           "nota": "Metricas de texto con Georgia (sistema) y DejaVu Sans como respaldo de Calibri (mas ancha: conservador)."}
    RUTA_JSON.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


if __name__ == "__main__":
    r = revisar()
    print(f"errores={r['errores']} avisos={r['avisos']}")
    for d in r["detalle"]:
        print(f"  [{d['nivel']}] d{d['diapositiva']} {d['regla']}: {d['mensaje']}")
    print("palabras:", r["palabras_por_diapositiva"])
